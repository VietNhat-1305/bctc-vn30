#!/usr/bin/env python3
"""Từ điển 500 thuật ngữ kinh tế - tài chính - kế toán - ngân hàng - chứng khoán.

Nguồn định nghĩa:
  1. Định nghĩa biên soạn (tu_dien/nguon/*.json): ngắn gọn, gắn bối cảnh BCTC Việt Nam, có công thức với chỉ số.
  2. Wikipedia tiếng Việt: đoạn mở đầu bài viết (API chính thức, theo trang chuyển hướng, bỏ trang định hướng),
     ghi nguồn + giấy phép CC BY-SA 4.0 như yêu cầu của Wikipedia.

    python tu_dien.py            # gộp nguồn + tra Wikipedia (có cache) -> tu_dien/tu_dien.json + tu_dien_500.xlsx
    python tu_dien.py --lam-moi  # tra lại Wikipedia, bỏ cache
"""
from __future__ import annotations

import json
import re
import sys
import time
import unicodedata
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parent
DIR = ROOT / "tu_dien"
API = "https://vi.wikipedia.org/w/api.php"
API_EN = "https://en.wikipedia.org/w/api.php"
# Wikimedia yêu cầu User-Agent mô tả rõ công cụ + liên hệ
UA = "BCTC-VN30-glossary/1.0 (https://github.com/VietNhat-1305/bctc-vn30; educational use)"   # header phải là ASCII
NHOM_THU_TU = ["Kinh tế vĩ mô & chính sách", "Kế toán & Báo cáo tài chính", "Phân tích tài chính & Định giá",
               "Ngân hàng & Tín dụng", "Chứng khoán & Đầu tư"]
STOP = {"va", "cua", "cac", "trong", "theo", "tren", "duoc", "mot", "nhung", "voi", "cho", "tu", "la", "co", "khong"}
# Đã soát tay: Wikipedia trả về bài sai khái niệm (chữ cái Beta/G, giao thức FTP, Viện PMI...) hoặc quá chung chung
WIKI_BO_QUA = {"Hệ số beta", "Tốc độ tăng trưởng dài hạn", "Định giá điều chuyển vốn nội bộ", "Chỉ số nhà quản trị mua hàng",
               "Cho vay tiêu dùng", "Công ty tài chính", "Hệ thống KRX"}
# Bước tra tên tiếng Anh: bài sai nghĩa (VD "Rủi ro phi hệ thống" -> Systematic risk, "Giá trần" -> giá trần do nhà nước áp)
# hoặc quá chung (Tài sản, Nợ, Lãi suất...) - soát tay sau mỗi lần thêm thuật ngữ
# Rà soát 17/10/2026 (5 agent đọc đủ 500 từ + 1 agent mở trang thật 20 từ ngẫu nhiên, đã xét lại từng kết luận):
# chọn đúng bài cho thuật ngữ bị ghép sai (None = không có bài phù hợp, bỏ trích Wikipedia)
WIKI_CHON: dict[str, tuple[str, str] | None] = {
    "Biên lợi nhuận hoạt động": ("en", "Operating margin"),          # bài vi cũ định nghĩa biên lợi nhuận RÒNG
    "Biên lợi nhuận ròng": ("vi", "Tỷ suất lợi nhuận trên doanh thu"),
    "Nợ và Có": ("en", "Debits and credits"),                         # "Debit and Credit" là tiểu thuyết Đức 1855
    "Biên độ dao động giá": ("en", "Price limit"),                    # "Trading band" nói về biên độ tỷ giá, Bollinger
    "Tỷ suất sinh lời": ("en", "Rate of return"),                     # bài vi dịch máy, có câu vô nghĩa
    "Tỷ suất lợi nhuận trên tổng tài sản": ("en", "Return on assets"),  # bài vi ghi sai tên khác
    "Tỷ giá hối đoái thực": None,           # đoạn mở đầu bài vi định nghĩa tỷ giá danh nghĩa
    "Hợp nhất kinh doanh": None,            # "Consolidation (business)" dễ nhầm với lập BCTC hợp nhất
    "Ngày đăng ký cuối cùng": None,         # "Ex-dividend date" là ngày giao dịch không hưởng quyền, khác ngày ĐKCC
    "Giá tham chiếu": None,                 # "Reference price" nói về mua sắm, y tế
    "Thặng dư sản xuất": None,              # đoạn trích chỉ nói về thặng dư tiêu dùng
    # soát vòng 2: bài thay thế vẫn chưa đạt
    "Cho vay tiêu dùng": None,              # bài "Tín dụng tiêu dùng" không có câu định nghĩa, có câu lỗi thời, dở dang
    "Lợi nhuận gộp": None,                  # "Gross income" mở đầu bằng thu nhập gộp cá nhân, gọi sai là credit sales
}
# Đúng bài nhưng câu sau của đoạn trích sai/lỗi thời -> chỉ lấy n câu đầu
WIKI_SO_CAU = {"Nhập siêu": 2, "Cổ phiếu phổ thông": 2, "Cổ phiếu ESOP": 1, "Chỉ số giảm phát GDP": 2,
               "Báo cáo kết quả hoạt động kinh doanh": 1, "Giá trị còn lại": 2,
               "Biên lợi nhuận hoạt động": 1, "Tỷ suất lợi nhuận trên tổng tài sản": 1, "GDP danh nghĩa": 1,
               "Công ty tài chính": 3, "Lệnh thị trường": 3}
# Bài nói khái niệm rộng hơn / hẹp hơn: vẫn giúp hiểu nhưng ghi rõ "gần đúng" (Excel + web)
WIKI_GAN_DUNG = {
    "Chính sách tài khóa mở rộng", "Chính sách tài khóa thắt chặt", "Chính sách tiền tệ nới lỏng", "Chính sách tiền tệ thắt chặt",
    "Đầu tư công", "GDP bình quân đầu người", "GDP danh nghĩa", "GDP thực tế", "Giá cân bằng", "Lãi suất tái chiết khấu",
    "Nợ công", "Thặng dư tiêu dùng", "Thiểu phát", "Tổng thu nhập quốc gia", "Tỷ lệ thất nghiệp", "Xếp hạng tín nhiệm quốc gia",
    "Xuất siêu", "Chi phí trả trước", "Cơ sở dồn tích", "Công ty mẹ", "Hao mòn lũy kế", "Khấu hao đường thẳng", "Nguyên giá",
    "Sổ kế toán", "Tài khoản kế toán", "Tài sản cố định hữu hình", "Tài sản cố định thuê tài chính", "Biên EBITDA",
    "Hệ số đòn bẩy tài chính", "Mô hình Gordon", "Phân tích độ nhạy", "Phân tích kịch bản", "Phân tích tỷ số tài chính",
    "Phân tích xu hướng", "Thu nhập trên mỗi cổ phiếu", "Thu nhập trên mỗi cổ phiếu pha loãng", "Vốn lưu động",
    "Chấm điểm tín dụng", "Công ty tài chính", "Hiệp ước Basel", "Lãi suất liên ngân hàng", "Tiền gửi tiết kiệm",
    "Xóa nợ ngoại bảng", "Call margin", "Cổ tức bằng cổ phiếu", "Cổ tức bằng tiền mặt", "Lệnh giới hạn", "Lệnh thị trường",
    "Phát hành thêm cổ phiếu", "Quyền mua cổ phiếu", "Thanh khoản", "Thị trường giá lên", "Thị trường giá xuống",
    "Chứng chỉ quỹ", "Công ty chứng khoán", "Cho vay tiêu dùng",
}
WIKI_BO_QUA_EN = {"Doanh thu thuần", "Hệ số thanh toán tiền mặt", "Nguyên tắc thận trọng", "Khoản mục bất thường",
                  "Nợ vay ròng", "Khối ngoại", "Lãi suất cho vay", "Lãi suất huy động", "Thị trường liên ngân hàng",
                  "Rủi ro phi hệ thống", "Tỷ lệ khả năng chi trả", "Tỷ lệ nguồn vốn ổn định ròng", "Tài sản ngắn hạn",
                  "Tài sản dài hạn", "Thuyết minh báo cáo tài chính", "Lợi nhuận thuần từ hoạt động kinh doanh",
                  "Lợi nhuận trước thuế", "Bất động sản đầu tư", "Chi phí nợ", "Vốn tự có", "Lợi suất trái phiếu",
                  "Giá trần", "Giá sàn", "Phân bổ chi phí"}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower().replace("đ", "d"))
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", "".join(c for c in s if unicodedata.category(c) != "Mn"))).strip()


def load_sources() -> list[dict]:
    """Gộp các file nhóm, bỏ trùng theo tên chuẩn hoá (giữ bản xuất hiện trước)."""
    items, seen = [], set()
    for f in sorted((DIR / "nguon").glob("*.json")):
        for it in json.loads(f.read_text(encoding="utf-8")):
            k = norm(it["thuat_ngu"])
            if not k or k in seen:
                continue
            seen.add(k)
            items.append({"thuat_ngu": it["thuat_ngu"].strip(), "viet_tat": (it.get("viet_tat") or "").strip() or None,
                          "tieng_anh": (it.get("tieng_anh") or "").strip(), "nhom": it.get("nhom", "").strip(),
                          "dinh_nghia": it["dinh_nghia"].strip(), "lien_quan": it.get("lien_quan") or []})
    return items


ABBR = re.compile(r"(?:\b[A-ZĐ]|\b(?:Mr|Mrs|Ms|Dr|St|Jr|Sr|Inc|Co|Ltd|Corp|vs|TP|Tp|No|U\.S|e\.g|i\.e))$")


def sentence_ends(text: str) -> list[int]:
    """Vị trí kết thúc câu - bỏ qua chữ viết tắt/tên viết tắt ("Myron J. Gordon", "U.S.", "TP. HCM")."""
    return [m.end() for m in re.finditer(r"[.!?](?=\s+\(?[A-ZĐÀ-Ỹ0-9]|$)", text) if not ABBR.search(text[:m.start()])]


def strip_math(text: str) -> str:
    """Bỏ khối công thức LaTeX "{\\displaystyle ...}" mà bản chữ thuần của Wikipedia để lại (đếm ngoặc lồng nhau)."""
    out, i = [], 0
    while True:
        j = text.find("{\\displaystyle", i)
        if j < 0:
            return "".join(out) + text[i:]
        out.append(text[i:j])
        depth, k = 0, j
        while k < len(text):
            depth += {"{": 1, "}": -1}.get(text[k], 0)
            k += 1
            if depth == 0:
                break
        i = k


def first_sentences(text: str, limit: int = 700, so_cau: int | None = None) -> str:
    """Đoạn mở đầu Wikipedia, cắt ở ranh giới câu, tối đa ~limit ký tự (hoặc so_cau câu đầu)."""
    text = re.sub(r"\s+", " ", re.sub(r"\(\s*[;,]?\s*\)", "", strip_math(text))).strip()
    text = re.sub(r"\s+([.,;:])", r"\1", text)
    ends = sentence_ends(text)
    if so_cau and len(ends) >= so_cau and len(text) > ends[so_cau - 1]:
        text = text[:ends[so_cau - 1]]
        ends = ends[:so_cau]
    if len(text) <= limit:
        return text
    ok = [e for e in ends if e <= limit]
    return text[:ok[-1]] if ok else text[:limit].rsplit(" ", 1)[0] + "…"


def en_names(it: dict) -> list[str]:
    """Tên tiếng Anh để tra en.wikipedia: tách 'A / B', thêm bản bỏ ngoặc và phần trong ngoặc."""
    out = []
    for part in re.split(r"\s+/\s+|;\s*", it["tieng_anh"]):
        for p in (part, re.sub(r"\s*\(.*?\)", "", part), *re.findall(r"\((.*?)\)", part)):
            p = p.strip()
            if p and p not in out:
                out.append(p)
    return out


def wiki_lookup(titles: list[str], client: httpx.Client, api: str = API, langlinks: bool = False) -> dict[str, dict]:
    """{tiêu đề hỏi: {title, url, extract[, vi_title]}} cho các bài tồn tại, không phải trang định hướng. 20 tiêu đề/lần.
    langlinks=True (dùng với en.wikipedia): kèm tên bài tiếng Việt tương ứng nếu có."""
    out = {}
    for i in range(0, len(titles), 20):
        chunk = titles[i:i + 20]
        params = {"action": "query", "format": "json", "formatversion": "2", "redirects": "1",
                  "prop": "extracts|pageprops|info", "inprop": "url", "exintro": "1", "explaintext": "1",
                  "exlimit": "20", "ppprop": "disambiguation", "titles": "|".join(chunk)}
        if langlinks:
            params.update(prop=params["prop"] + "|langlinks", lllang="vi", lllimit="500")
        for attempt in range(4):
            try:
                r = client.get(api, params=params, timeout=30)
                r.raise_for_status()
                data = r.json()["query"]
                break
            except Exception:  # noqa: BLE001
                if attempt == 3:
                    raise
                time.sleep(2 ** attempt)
        # hỏi -> (chuẩn hoá chữ hoa đầu) -> (chuyển hướng) -> tiêu đề bài thật
        mapping = {t: t for t in chunk}
        for n in data.get("normalized", []):
            for k, v in mapping.items():
                if v == n["from"]:
                    mapping[k] = n["to"]
        for rd in data.get("redirects", []):
            for k, v in mapping.items():
                if v == rd["from"]:
                    mapping[k] = rd["to"]
        pages = {p["title"]: p for p in data.get("pages", [])}
        for asked, real in mapping.items():
            p = pages.get(real)
            if not p or p.get("missing") or "disambiguation" in (p.get("pageprops") or {}):
                continue
            ext = (p.get("extract") or "").strip()
            if len(ext) >= 60:
                out[asked] = {"title": p["title"], "url": p["fullurl"], "extract": first_sentences(ext)}
                vi = [x["title"] for x in p.get("langlinks") or [] if x.get("lang") == "vi"]
                if vi:
                    out[asked]["vi_title"] = vi[0]
        time.sleep(0.3)
    return out


def relevant(item: dict, wiki: dict) -> bool:
    """Bài tìm qua TỪ VIẾT TẮT (VD 'NIM', 'CAR') phải nói đúng khái niệm: có từ khoá của tên/tiếng Anh."""
    words = {w for w in norm(item["thuat_ngu"] + " " + item["tieng_anh"]).split() if len(w) > 2 and w not in STOP}
    text = norm(wiki["title"] + " " + wiki["extract"])
    return sum(w in text for w in words) >= max(1, len(words) // 3)


def enrich(items: list[dict], refresh: bool = False) -> None:
    cache_path = DIR / "wiki_cache.json"
    cache = {} if refresh or not cache_path.exists() else json.loads(cache_path.read_text(encoding="utf-8"))
    asks = []
    for it in items:
        for t in filter(None, [it["thuat_ngu"], it["viet_tat"]]):
            if t not in cache:
                asks.append(t)
    asks = list(dict.fromkeys(asks))
    if asks:
        print(f"Tra Wikipedia tiếng Việt: {len(asks)} tiêu đề ({(len(asks) + 19) // 20} lần gọi API)...")
        with httpx.Client(headers={"User-Agent": UA}, follow_redirects=True) as client:
            found = wiki_lookup(asks, client)
        for t in asks:
            cache[t] = found.get(t)                      # None = không có bài phù hợp (cũng cache lại)
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    for it in items:
        w = cache.get(it["thuat_ngu"])
        if not w and it["viet_tat"]:
            w = cache.get(it["viet_tat"])
            if w and not relevant(it, w):
                w = None
        it["wiki"] = None if it["thuat_ngu"] in WIKI_BO_QUA or not w else {**w, "lang": "vi"}

    # Bước 2: chưa có bài tiếng Việt -> tra tên tiếng Anh trên en.wikipedia (tên chuẩn quốc tế, khớp chính xác hơn);
    # bài tiếng Anh có bản tiếng Việt thì lấy bản tiếng Việt, không thì trích bài tiếng Anh.
    need = [it for it in items if not it["wiki"] and it["thuat_ngu"] not in WIKI_BO_QUA_EN]
    asks = list(dict.fromkeys(n for it in need for n in en_names(it) if "en:" + n not in cache))
    if asks:
        print(f"Tra Wikipedia tiếng Anh: {len(asks)} tiêu đề ({(len(asks) + 19) // 20} lần gọi API)...")
        with httpx.Client(headers={"User-Agent": UA}, follow_redirects=True) as client:
            found = wiki_lookup(asks, client, api=API_EN, langlinks=True)
            for t in asks:
                cache["en:" + t] = found.get(t)
            vi_asks = list(dict.fromkeys(w["vi_title"] for w in found.values()
                                         if w.get("vi_title") and w["vi_title"] not in cache))
            if vi_asks:
                found_vi = wiki_lookup(vi_asks, client)
                for t in vi_asks:
                    cache[t] = found_vi.get(t)
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    for it in need:
        for n in en_names(it):
            w = cache.get("en:" + n)
            if w:
                vi = cache.get(w.get("vi_title") or "")
                it["wiki"] = {**vi, "lang": "vi"} if vi else {**w, "lang": "en"}
                break

    # Bước 3: kết quả rà soát thủ công - bài chọn tay, cắt câu sai, đánh dấu mức khớp
    picks = sorted({v for v in WIKI_CHON.values() if v and f"{v[0]}|{v[1]}" not in cache})
    if picks:
        with httpx.Client(headers={"User-Agent": UA}, follow_redirects=True) as client:
            for lang in ("vi", "en"):
                titles = [t for lg, t in picks if lg == lang]
                if titles:
                    found = wiki_lookup(titles, client, api=API if lang == "vi" else API_EN)
                    for t in titles:
                        cache[f"{lang}|{t}"] = found.get(t)
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    for it in items:
        name = it["thuat_ngu"]
        if name in WIKI_CHON:
            spec = WIKI_CHON[name]
            page = cache.get(f"{spec[0]}|{spec[1]}") if spec else None
            it["wiki"] = {**page, "lang": spec[0]} if page else None
        if it["wiki"]:
            if name in WIKI_SO_CAU:
                it["wiki"] = {**it["wiki"], "extract": first_sentences(it["wiki"]["extract"], so_cau=WIKI_SO_CAU[name])}
            it["wiki"]["muc_khop"] = "gần đúng" if name in WIKI_GAN_DUNG else "đúng"


def link_related(items: list[dict]) -> None:
    """Chỉ giữ thuật ngữ liên quan có trong từ điển (để bấm được), khớp không phân biệt dấu/hoa thường/viết tắt."""
    idx = {}
    for it in items:
        idx[norm(it["thuat_ngu"])] = it["thuat_ngu"]
        if it["viet_tat"]:
            idx.setdefault(norm(it["viet_tat"]), it["thuat_ngu"])
    for it in items:
        rel = []
        for r in it["lien_quan"]:
            t = idx.get(norm(r))
            if t and t != it["thuat_ngu"] and t not in rel:
                rel.append(t)
        it["lien_quan"] = rel[:6]


def write_excel(items: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Từ điển"
    head = ["STT", "Thuật ngữ", "Viết tắt", "Tiếng Anh", "Nhóm", "Định nghĩa (biên soạn)",
            "Trích Wikipedia", "Bài Wikipedia", "Ngôn ngữ bài", "Mức khớp Wikipedia", "Thuật ngữ liên quan"]
    ws.append(head)
    for i, it in enumerate(items, 1):
        w = it.get("wiki") or {}
        ws.append([i, it["thuat_ngu"], it["viet_tat"] or "", it["tieng_anh"], it["nhom"], it["dinh_nghia"],
                   w.get("extract", ""), w.get("url", ""), {"vi": "Tiếng Việt", "en": "Tiếng Anh"}.get(w.get("lang"), ""), w.get("muc_khop", ""),
                   ", ".join(it["lien_quan"])])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (6, 7, 11))
        if row[7].value:
            row[7].hyperlink = row[7].value
            row[7].font = Font(color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHIJK", (5, 30, 10, 28, 24, 70, 70, 34, 12, 14, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("Nguồn & giấy phép")
    for line in ["Định nghĩa (biên soạn): viết cho dự án BCTC VN30 (https://github.com/VietNhat-1305/bctc-vn30).",
                 "Cột 'Trích Wikipedia': đoạn mở đầu bài viết tương ứng trên vi.wikipedia.org (không có thì en.wikipedia.org, nguyên văn),",
                 "giấy phép Creative Commons Ghi công - Chia sẻ tương tự 4.0 (CC BY-SA 4.0):",
                 "https://creativecommons.org/licenses/by-sa/4.0/deed.vi - xem tác giả trong lịch sử từng bài (cột 'Bài Wikipedia')."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120
    wb.save(path)


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    items = load_sources()
    items.sort(key=lambda it: (NHOM_THU_TU.index(it["nhom"]) if it["nhom"] in NHOM_THU_TU else 9, norm(it["thuat_ngu"])))
    enrich(items, refresh="--lam-moi" in sys.argv)
    link_related(items)
    (DIR / "tu_dien.json").write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
    write_excel(items, ROOT / "tu_dien_500.xlsx")
    n_wiki = sum(1 for it in items if it.get("wiki"))
    by = {}
    for it in items:
        by[it["nhom"]] = by.get(it["nhom"], 0) + 1
    print(f"{len(items)} thuật ngữ | có bài Wikipedia: {n_wiki} | theo nhóm: {by}")
    print(f"-> {DIR / 'tu_dien.json'}  +  {ROOT / 'tu_dien_500.xlsx'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
