#!/usr/bin/env python3
"""Kiểm tra web IR của 30 công ty: đọc được danh sách BCTC? có BCTC quý mới nhất? tải được file?
So với CafeF để thấy web công ty có bản mới sớm hơn không. Xuất website_30cty.xlsx.

    python kiem_tra_ir.py                  # 30 mã
    python kiem_tra_ir.py --tickers VHM,SHB
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
import time
from datetime import date
from pathlib import Path
from urllib.parse import unquote, urlsplit

import aiohttp

import bctc

ROOT = Path(__file__).resolve().parent
QUARTERS = ("Q1", "Q2", "Q3", "Q4")
MODE_VI = {"static_html": "Đọc HTML", "json_api": "Gọi API JSON", "two_step": "2 bước (danh sách → bài)",
           "custom": "Adapter riêng"}
FETCH_VI = {"http": "", "curl": " qua curl", "browser": " qua Chrome headless"}


def ir_date(d):
    """Ngày công bố trên web công ty: tiền tố YYYYMMDD trong tên file/tiêu đề ('20260730 - VNM - ...')."""
    import re
    from datetime import datetime as dt
    m = re.search(r"(?<!\d)(20[2-4]\d)(0[1-9]|1[0-2])([0-3]\d)(?!\d)", f"{unquote(d.url)} {d.title}")
    try:
        return dt(int(m[1]), int(m[2]), int(m[3])) if m else None
    except ValueError:
        return None


def cafef_time(d):
    """Thời điểm CafeF đưa file lên: hậu tố _DDMMYYYYHHMMSS, _<unix giây> hoặc ?v=<unix mili giây>."""
    import re
    from datetime import datetime as dt
    u = d.url
    m = re.search(r"_(\d{2})(\d{2})(20\d{2})(\d{2})(\d{2})(\d{2})\.pdf", u, re.I)
    if m:
        return dt(int(m[3]), int(m[2]), int(m[1]), int(m[4]), int(m[5]))
    m = re.search(r"\?v=(1\d{12})$", u) or re.search(r"_(1\d{9})\.pdf", u, re.I)
    if m:
        ts = int(m[1])
        return dt.fromtimestamp(ts / 1000 if ts > 1e12 else ts)
    return None


def latest(docs, scopes=("HN", "CHUNG")):
    q = [d for d in docs if d.period in QUARTERS and d.scope in scopes and d.year]
    return max(q, key=lambda d: (d.year, int(d.period[1])), default=None)


def label(d) -> str:
    return f"{d.period}/{d.year}" if d else "-"


async def probe_file(session, co, url) -> str:
    """Tải thử 1 KB đầu: phải là PDF/ZIP/RAR... (không phải trang lỗi HTML)."""
    host = urlsplit(url).hostname or ""
    if co.get("fetcher") == "curl" and host in (urlsplit(co.get("bctc_url", "")).hostname or "",):
        tmp = ROOT / f".probe_{co['ticker']}.bin"
        status, _ = await bctc.curl_get(url, dest=tmp)
        head = tmp.read_bytes()[:16] if tmp.exists() else b""
        tmp.unlink(missing_ok=True)
    else:
        headers = {"Range": "bytes=0-1023"}
        if co.get("ua"):
            headers["User-Agent"] = co["ua"]
        async with session.get(url, headers=headers) as r:
            status = r.status
            head = await r.content.read(16)
    if status >= 400:
        return f"HTTP {status}"
    return "OK" if head.lstrip().startswith(bctc.MAGIC) else f"không phải PDF ({head[:8]!r})"


async def check(session, co, years) -> dict:
    t0 = time.perf_counter()
    res = {"co": co, "docs": 0, "ir": None, "cafef": None, "q_prev": False, "file": "", "status": "FAIL", "err": "",
           "lag": None}
    cafef = []
    try:
        cafef = await bctc.list_cafef(session, co["ticker"], 2)
        res["cafef"] = latest(cafef)
    except Exception as e:  # noqa: BLE001
        res["err"] += f"CafeF lỗi {e}; "
    try:
        docs = await bctc.list_official(session, co, 2, years)
        res["docs"] = len([d for d in docs if d.year in years])
        top = latest(docs)
        res["ir"] = top
        if top:
            prev = (top.year, int(top.period[1]) - 1) if top.period != "Q1" else (top.year - 1, 4)
            res["q_prev"] = any((d.year, int(d.period[1])) == prev for d in docs
                                if d.period in QUARTERS and d.scope in ("HN", "CHUNG"))
            res["file"] = top.url
            # CafeF đăng sau web công ty bao lâu (cùng kỳ, cùng phạm vi): ngày công bố trên web vs giờ CafeF upload
            dates = [x for x in (ir_date(d) for d in docs if d.key[:4] == top.key[:4]) if x]
            times = [x for x in (cafef_time(d) for d in cafef if d.key[:4] == top.key[:4]) if x]
            if dates and times:
                res["lag"] = (min(times) - min(dates)).total_seconds() / 86400
            res["probe"] = await probe_file(session, co, top.url)
            ok_period = res["cafef"] is None or (top.year, top.period) >= (res["cafef"].year, res["cafef"].period)
            if res["probe"] == "OK" and ok_period:
                res["status"] = "PASS"
            else:
                res["err"] += f"file: {res['probe']}; " if res["probe"] != "OK" else "web IR cũ hơn CafeF; "
        else:
            res["err"] += "không thấy BCTC quý hợp nhất; "
    except Exception as e:  # noqa: BLE001
        res["err"] += f"{type(e).__name__}: {str(e)[:120]}"
    res["sec"] = time.perf_counter() - t0
    return res


def write_json(rows: list[dict], path: Path):
    """Kết quả kiểm tra cho web (dashboard.py): BCTC quý mới nhất trên web công ty, độ trễ CafeF..."""
    from datetime import datetime as dt

    def doc(d):
        return None if not d else {"ky": label(d), "pham_vi": d.scope, "tieu_de": d.title, "url": d.url,
                                   "ngay": (ir_date(d).strftime("%Y-%m-%d") if ir_date(d) else None)}
    out = [{"ma": r["co"]["ticker"], "ten": r["co"].get("name", ""), "website": r["co"].get("website", ""),
            "bctc_url": r["co"].get("bctc_url", ""),
            "cach_doc": MODE_VI.get(r["co"].get("fetch_mode"), "") + FETCH_VI.get(r["co"].get("fetcher", "http"), ""),
            "kiem_tra": r["status"], "ir": doc(r["ir"]), "cafef_ky": label(r["cafef"]),
            "cafef_tre_ngay": round(r["lag"], 1) if r["lag"] is not None else None,
            "co_quy_truoc": r["q_prev"], "so_bctc": r["docs"], "giay": round(r["sec"], 1), "loi": r["err"].strip("; ")}
           for r in rows]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"kiem_tra_luc": dt.now().isoformat(timespec="minutes"), "cong_ty": out},
                               ensure_ascii=False, indent=1), encoding="utf-8")


def write_excel(rows: list[dict], path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Website 30 công ty"
    head = ["STT", "Mã", "Tên công ty", "Website", "Trang đăng BCTC", "Trang BCTC (EN)", "Cách tool đọc",
            "Kiểm tra", "BCTC quý mới nhất trên web công ty", "Mới nhất trên CafeF", "Web công ty sớm hơn?",
            "Có quý liền trước?", "CafeF đăng sau web công ty (ngày)", "Số BCTC 2025-2026 tìm thấy", "Link BCTC mới nhất", "Thời gian đọc (giây)",
            "Ghi chú kỹ thuật", "Lỗi"]
    ws.append(head)
    for i, r in enumerate(rows, 1):
        co = r["co"]
        ir, cf = r["ir"], r["cafef"]
        earlier = ("CÓ" if ir and cf and (ir.year, ir.period) > (cf.year, cf.period) else
                   "Bằng" if ir and cf and (ir.year, ir.period) == (cf.year, cf.period) else "-")
        ws.append([i, co["ticker"], co.get("name", ""), co.get("website", ""), co.get("bctc_url", ""),
                   co.get("bctc_url_en") or "",
                   MODE_VI.get(co.get("fetch_mode"), co.get("fetch_mode")) + FETCH_VI.get(co.get("fetcher", "http"), ""),
                   r["status"], label(ir), label(cf), earlier, "Có" if r["q_prev"] else "Không",
                   round(r["lag"], 1) if r["lag"] is not None else "", r["docs"],
                   r["file"], round(r["sec"], 1), co.get("note", ""), r["err"].strip("; ")])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for row in ws.iter_rows(min_row=2):
        row[7].fill = green if row[7].value == "PASS" else red
        for col in (3, 4, 5, 14):
            if row[col].value:
                row[col].hyperlink = row[col].value
                row[col].font = Font(color="0563C1", underline="single")
    widths = [5, 7, 40, 26, 45, 30, 26, 9, 16, 12, 12, 11, 14, 12, 50, 10, 60, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


async def run(args) -> int:
    data = json.loads((ROOT / "companies.json").read_text(encoding="utf-8"))
    cos = [c for c in data["companies"] if not args.tickers or c["ticker"] in args.tickers]
    this = date.today().year
    years = {this - 1, this}
    t0 = time.perf_counter()
    conn = aiohttp.TCPConnector(limit=24, limit_per_host=4)
    timeout = aiohttp.ClientTimeout(total=None, connect=20, sock_read=60)
    async with aiohttp.ClientSession(connector=conn, timeout=timeout,
                                     headers={"User-Agent": bctc.UA, "Accept-Language": "vi,en;q=0.8"}) as session:
        rows = await asyncio.gather(*(check(session, c, years) for c in cos))
    rows.sort(key=lambda r: r["co"]["ticker"])
    print(f"\n{'Mã':<5}{'Kết quả':<8}{'Web IR':>9}{'CafeF':>9}  {'Q trước':<8}{'CafeF trễ':>12}{'Số BCTC':>8}{'Giây':>7}  Ghi chú")
    for r in rows:
        lag = f"{r['lag']:+.1f} ngày" if r["lag"] is not None else "-"
        print(f"{r['co']['ticker']:<5}{r['status']:<8}{label(r['ir']):>9}{label(r['cafef']):>9}  "
              f"{'có' if r['q_prev'] else 'không':<8}{lag:>12}{r['docs']:>8}{r['sec']:>7.1f}  {r['err'].strip('; ')[:70]}")
    n_pass = sum(r["status"] == "PASS" for r in rows)
    lags = sorted(r["lag"] for r in rows if r["lag"] is not None)
    if lags:
        print(f"\nCafeF đăng sau web công ty (BCTC quý mới nhất, {len(lags)} mã đo được): trung vị "
              f"{lags[len(lags) // 2]:.1f} ngày, chậm nhất {lags[-1]:.1f} ngày")
    early = [r["co"]["ticker"] for r in rows if r["ir"] and r["cafef"]
             and (r["ir"].year, r["ir"].period) > (r["cafef"].year, r["cafef"].period)]
    out = ROOT / ("website_30cty.xlsx" if not args.tickers else "website_kiemtra_tam.xlsx")   # chạy 1 phần thì không ghi đè file 30 mã
    write_excel(rows, out)
    if not args.tickers:
        write_json(rows, ROOT / "ket_qua" / "ir_check.json")
    print(f"\nPASS {n_pass}/{len(rows)} trong {time.perf_counter() - t0:.0f}s | web công ty có bản mới hơn CafeF: "
          f"{', '.join(early) or 'không'} | -> {out}")
    return 0 if n_pass == len(rows) else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Kiểm tra đọc BCTC từ web IR của từng công ty.")
    ap.add_argument("--tickers", type=lambda s: [x.strip().upper() for x in s.split(",") if x.strip()])
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    return asyncio.run(run(args))


if __name__ == "__main__":
    sys.exit(main())
